#!/usr/bin/env python3
"""Build a weighted 6-criteria mining scorecard (.xlsx) from a data.json.

Usage:  python build_scorecard.py data.json output.xlsx

The scorecard uses SUMPRODUCT + RANK formulas (results are NOT hardcoded), so the
user can change any score or weight and the totals/ranking update automatically.
Requires openpyxl (preinstalled). Run the xlsx recalc afterwards if you want cached
values, e.g. LibreOffice headless convert, but formulas are valid as written.
"""
import json, sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter

CRITERIA = [
    ("reserves",   "Trữ lượng & tuổi thọ mỏ",        0.25),
    ("production", "Sản lượng khai thác",            0.10),
    ("margin",     "Biên lợi nhuận",                 0.20),
    ("fcf",        "Dòng tiền tự do (FCF)",          0.20),
    ("debt",       "Nợ & bảng cân đối",              0.15),
    ("valuation",  "Định giá (EV/EBITDA, P/NAV)",    0.10),
]
ARIAL = "Arial"
NAVY = "1F3864"; BLUE = "2E5496"; GREY = "D9D9D9"; YELLOW = "FFF2CC"


def main(data_path, out_path):
    with open(data_path, encoding="utf-8") as f:
        d = json.load(f)
    companies = d["companies"]
    weights = d.get("weights") or {k: w for k, _, w in CRITERIA}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bảng chấm điểm"

    title_font = Font(name=ARIAL, size=14, bold=True, color="FFFFFF")
    hdr_font = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
    crit_font = Font(name=ARIAL, size=10)
    input_font = Font(name=ARIAL, size=10, color="0000FF")
    formula_font = Font(name=ARIAL, size=10, bold=True)
    navy = PatternFill("solid", fgColor=NAVY)
    blue = PatternFill("solid", fgColor=BLUE)
    grey = PatternFill("solid", fgColor=GREY)
    yellow = PatternFill("solid", fgColor=YELLOW)
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    tickers = [c["ticker"] for c in companies]
    ncol = 2 + len(tickers)
    last_col = get_column_letter(ncol)

    ws.merge_cells(f"A1:{last_col}1")
    ws["A1"] = "BẢNG CHẤM ĐIỂM MÔ HÌNH KHOÁNG SẢN — " + " · ".join(tickers)
    ws["A1"].font = title_font; ws["A1"].fill = navy; ws["A1"].alignment = center
    ws.row_dimensions[1].height = 26
    ws.merge_cells(f"A2:{last_col}2")
    ws["A2"] = (f"Thang điểm 1–5 (5 = tốt nhất), có trọng số · {d.get('report_date','')} · "
                f"Số liệu tại {d.get('data_date','')} · Ô xanh = có thể điều chỉnh")
    ws["A2"].font = Font(name=ARIAL, size=9, italic=True, color="555555"); ws["A2"].alignment = left

    hr = 4
    ws.cell(hr, 1, "Tiêu chí").font = hdr_font
    ws.cell(hr, 2, "Trọng số").font = hdr_font
    for j, t in enumerate(tickers):
        ws.cell(hr, 3 + j, t).font = hdr_font
    for col in range(1, ncol + 1):
        c = ws.cell(hr, col); c.fill = blue; c.alignment = center; c.border = border
    ws.row_dimensions[hr].height = 20

    first = hr + 1
    for i, (key, label, _) in enumerate(CRITERIA):
        r = first + i
        a = ws.cell(r, 1, label); a.font = crit_font; a.alignment = left; a.border = border
        b = ws.cell(r, 2, weights.get(key, 0)); b.font = input_font; b.number_format = "0%"
        b.alignment = center; b.border = border; b.fill = yellow
        for j, comp in enumerate(companies):
            sc = comp.get("scores", {}).get(key)
            cell = ws.cell(r, 3 + j, sc); cell.font = input_font; cell.alignment = center; cell.border = border
    last = first + len(CRITERIA) - 1

    wc = last + 1
    ws.cell(wc, 1, "Tổng trọng số (phải = 100%)").font = Font(name=ARIAL, size=9, italic=True)
    chk = ws.cell(wc, 2, f"=SUM(B{first}:B{last})")
    chk.number_format = "0%"; chk.font = Font(name=ARIAL, size=9, italic=True, bold=True); chk.alignment = center

    tr = wc + 1
    ws.cell(tr, 1, "ĐIỂM TỔNG (có trọng số)").font = formula_font
    ws.cell(tr, 1).fill = grey; ws.cell(tr, 1).border = border
    ws.cell(tr, 2, "").fill = grey
    for j in range(len(companies)):
        col = get_column_letter(3 + j)
        cell = ws.cell(tr, 3 + j, f"=SUMPRODUCT($B${first}:$B${last},{col}{first}:{col}{last})")
        cell.font = formula_font; cell.number_format = "0.00"; cell.alignment = center
        cell.fill = grey; cell.border = border

    rr = tr + 1
    ws.cell(rr, 1, "XẾP HẠNG").font = formula_font; ws.cell(rr, 1).border = border
    f_first = get_column_letter(3); f_last = get_column_letter(2 + len(companies))
    for j in range(len(companies)):
        col = get_column_letter(3 + j)
        cell = ws.cell(rr, 3 + j, f"=RANK({col}{tr},${f_first}${tr}:${f_last}${tr},0)")
        cell.font = formula_font; cell.alignment = center; cell.border = border

    cs = ColorScaleRule(start_type="num", start_value=1, start_color="F8696B",
                        mid_type="num", mid_value=3, mid_color="FFEB84",
                        end_type="num", end_value=5, end_color="63BE7B")
    ws.conditional_formatting.add(f"C{first}:{last_col}{last}", cs)

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 10
    for j in range(len(companies)):
        ws.column_dimensions[get_column_letter(3 + j)].width = 9

    lr = rr + 2
    ws.cell(lr, 1, "Thang điểm:").font = Font(name=ARIAL, size=9, bold=True)
    for k, t in enumerate(["1 = Rất yếu", "2 = Yếu", "3 = Trung bình", "4 = Tốt", "5 = Rất tốt"]):
        ws.cell(lr + 1 + k, 1, t).font = Font(name=ARIAL, size=9)

    # ---- Data sheet ----
    ds = wb.create_sheet("Dữ liệu nền")
    ds.merge_cells("A1:H1")
    ds["A1"] = f"DỮ LIỆU NỀN (tại {d.get('data_date','')})"
    ds["A1"].font = Font(name=ARIAL, size=12, bold=True, color="FFFFFF"); ds["A1"].fill = navy
    ds["A1"].alignment = center; ds.row_dimensions[1].height = 22
    fields = [("Mã", "ticker"), ("Khoáng sản", "commodity"), ("Trữ lượng & đời mỏ", "reserves"),
              ("Sản lượng/năm", "production"), ("Biên LN ròng", "net_margin"),
              ("FCF & Nợ", "fcf_debt"), ("ROE", "roe"), ("P/E · EV/EBITDA", None)]
    for j, (h, _) in enumerate(fields):
        c = ds.cell(3, j + 1, h); c.font = hdr_font; c.fill = blue; c.alignment = center; c.border = border
    for i, comp in enumerate(companies):
        for j, (h, key) in enumerate(fields):
            if key is None:
                val = f"{comp.get('pe','')} · {comp.get('ev_ebitda','')}"
            else:
                val = comp.get(key, "")
            cell = ds.cell(4 + i, j + 1, val)
            cell.font = Font(name=ARIAL, size=9, bold=(j == 0)); cell.alignment = left; cell.border = border
    ds.column_dimensions["A"].width = 8
    for j, (h, _) in enumerate(fields[1:], start=2):
        ds.column_dimensions[get_column_letter(j)].width = 22
    nr = 4 + len(companies) + 1
    ds.cell(nr, 1, "Nguồn: " + d.get("sources", "")).font = Font(name=ARIAL, size=8, italic=True, color="555555")

    wb.save(out_path)
    print(f"saved {out_path}")


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python build_scorecard.py data.json output.xlsx"); sys.exit(1)
    main(sys.argv[1], sys.argv[2])
